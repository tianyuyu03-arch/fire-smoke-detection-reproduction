from __future__ import annotations

import sys
import tempfile
import subprocess
from pathlib import Path


PROJECT_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = PROJECT_DIR.parent
BUILD_DIR = PROJECT_DIR / ".build"
PDF_PATH = PROJECT_DIR / "main.pdf"
PPTX_PATH = PROJECT_DIR / "AEL-YOLO_组会汇报.pptx"

# Keep the presentation dependency local to this project.
sys.path.insert(0, str(BUILD_DIR / "pptx-python"))
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402
from PIL import Image  # noqa: E402


def main() -> None:
    if not PDF_PATH.is_file():
        raise FileNotFoundError(f"PDF not found: {PDF_PATH}")

    BUILD_DIR.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="pptx-render-", dir=BUILD_DIR) as temp_name:
        temp_dir = Path(temp_name)
        prefix = temp_dir / "slide"
        subprocess.run(
            [
                "pdftoppm",
                "-png",
                "-scale-to",
                "1920",
                str(PDF_PATH),
                str(prefix),
            ],
            check=True,
        )
        rendered_pages = sorted(temp_dir.glob("slide-*.png"))
        if not rendered_pages:
            raise RuntimeError("PDF rendering produced no slide images")
        page_images = []
        for rendered_path in rendered_pages:
            image_path = rendered_path.with_suffix(".jpg")
            with Image.open(rendered_path) as page_image:
                page_image.convert("RGB").save(image_path, "JPEG", quality=95, optimize=True)
            page_images.append(image_path)

        presentation = Presentation()
        presentation.slide_width = Inches(13.333333)
        presentation.slide_height = Inches(7.5)
        presentation.core_properties.title = "APConv / AEL-YOLO 组会汇报（图片版）"
        presentation.core_properties.subject = "每页以整页图片呈现，保持与 PDF 版式一致"

        blank_layout = presentation.slide_layouts[6]
        for image_path in page_images:
            slide = presentation.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                str(image_path),
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height,
            )

        presentation.save(PPTX_PATH)
        print(f"Created {PPTX_PATH} ({len(page_images)} image-only slides)")


if __name__ == "__main__":
    main()

