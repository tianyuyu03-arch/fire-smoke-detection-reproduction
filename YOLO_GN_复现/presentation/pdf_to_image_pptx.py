from __future__ import annotations

import subprocess
import shutil
import sys
import tempfile
from pathlib import Path

PRESENTATION_DIR = Path(__file__).resolve().parent
PROJECT_DIR = PRESENTATION_DIR.parent
BUILD_DIR = PRESENTATION_DIR / ".build"
sys.path.insert(0, str(BUILD_DIR / "pptx-python"))

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

BASE_NAME = "YOLO_GN_" + "".join(map(chr, (0x8bba, 0x6587, 0x7cbe, 0x8bfb, 0x4e0e, 0x590d, 0x73b0, 0x5f, 0x7ec4, 0x4f1a, 0x6c47, 0x62a5)))
PDF_PATH = PROJECT_DIR / f"{BASE_NAME}.pdf"
PPTX_PATH = PROJECT_DIR / f"{BASE_NAME}.pptx"


def main() -> None:
    built_pdf = BUILD_DIR / "YOLO_GN.pdf"
    if not PDF_PATH.is_file() and built_pdf.is_file():
        shutil.copy2(built_pdf, PDF_PATH)
    if not PDF_PATH.is_file():
        raise FileNotFoundError(f"Compile the PDF first: {PDF_PATH}")
    with tempfile.TemporaryDirectory(prefix="slides-", dir=BUILD_DIR) as temp_name:
        temp_dir = Path(temp_name)
        prefix = temp_dir / "slide"
        subprocess.run(
            ["pdftoppm", "-png", "-scale-to", "1920", str(PDF_PATH), str(prefix)],
            check=True,
        )
        rendered = sorted(temp_dir.glob("slide-*.png"))
        if not rendered:
            raise RuntimeError("PDF rendering produced no slide images")

        presentation = Presentation()
        presentation.slide_width = Inches(13.333333)
        presentation.slide_height = Inches(7.5)
        presentation.core_properties.title = "YOLO-GN+BIFPN 论文精读与复现结果"
        presentation.core_properties.subject = "静态图片版，与 PDF 页面一致"
        blank = presentation.slide_layouts[6]
        for image_path in rendered:
            slide = presentation.slides.add_slide(blank)
            slide.shapes.add_picture(
                str(image_path),
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height,
            )
        presentation.save(PPTX_PATH)
        print(f"Created {PPTX_PATH} ({len(rendered)} static slides)")


if __name__ == "__main__":
    main()

